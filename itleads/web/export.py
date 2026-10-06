"""The downloads: the same rows as the Google Sheet, as Excel or CSV."""
from __future__ import annotations

import csv
import hashlib
import io
import re
import threading
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import config, sheet

# header, key, width (characters)
COLUMNS = [
    ("Company", "company", 34), ("Website", "website", 26), ("Email", "email", 32), ("Phone", "phone", 17),
    ("Address", "address", 52), ("Registered", "registered", 13), ("Contact", "contact", 32),
    ("Contact email", "contact_email", 30), ("LinkedIn", "linkedin", 34), ("Industry", "industry", 32),
    ("State", "state", 7), ("Source", "source", 15), ("Fit", "fit", 12), ("Verified by", "proof", 70),
]
LINKS = {"website": lambda v: v, "email": lambda v: "mailto:" + v, "linkedin": lambda v: v}
ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")      # characters a spreadsheet file cannot hold
FORMULA = ("=", "+", "-", "@", "\t", "\r")
SAFE_LINK = re.compile(r"^(https?://[^\s]+|mailto:[^\s@]+@[^\s@]+)$", re.I)
_build_lock = threading.Lock()


def rows_for(store) -> list:
    return [sheet.sheet_row(it) for it in store.export_items()]


def _clean(v) -> str:
    return ILLEGAL.sub("", str(v or "").encode("utf-8", "ignore").decode("utf-8", "ignore"))


def _site(url: str) -> str:
    return url.replace("https://", "").replace("http://", "").removeprefix("www.").rstrip("/")


def csv_safe(v) -> str:
    """Spreadsheet programs run text that starts with = + - @ as a formula; keep every value inert."""
    v = _clean(v)
    return "'" + v if v[:1] in FORMULA else v


_csv_safe = csv_safe


def to_csv(rows: list) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for h, _, _ in COLUMNS])
    for r in rows:
        w.writerow([csv_safe(r.get(k, "")) for _, k, _ in COLUMNS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")           # the BOM lets Excel read UTF-8


def to_xlsx(rows: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    ws.sheet_view.showGridLines = False
    head_font = Font(name="Calibri", size=9, bold=True, color="6B6F6A")
    head_fill = PatternFill("solid", fgColor="F4F4F0")
    edge = Border(bottom=Side(style="medium", color="E3E4DF"))
    hair = Border(bottom=Side(style="thin", color="EEEEEA"))
    body = Font(name="Calibri", size=11, color="1D1F1E")
    link = Font(name="Calibri", size=11, color="14614F")
    for c, (h, _, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=h.upper())
        cell.font, cell.fill, cell.border = head_font, head_fill, edge
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 26
    for i, r in enumerate(rows, 2):
        ws.row_dimensions[i].height = 22
        for c, (_, key, _) in enumerate(COLUMNS, 1):
            v = _clean(r.get(key, ""))
            cell = ws.cell(row=i, column=c)
            cell.font, cell.border, cell.alignment = body, hair, Alignment(vertical="center")
            if key == "registered" and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
                y, m, d = (int(x) for x in v.split("-"))
                cell.value, cell.number_format = date(y, m, d), "d mmm yyyy"
                continue
            cell.value = _site(v) if key == "website" and v else v
            if isinstance(cell.value, str):
                cell.data_type = "s"                              # every value is text: nothing here is ever a formula
            if key in LINKS and v:
                target = LINKS[key](v)
                if SAFE_LINK.match(target):                       # only web and mail links, never javascript: or file:
                    cell.hyperlink = target
                    cell.font = link
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def cached(store, kind: str) -> bytes:
    """Excel files take seconds to build for a big list: build once per change of data, and let simultaneous
    requests share that one build instead of stalling the server."""
    version = hashlib.sha1(repr(store.export_version()).encode()).hexdigest()[:16]
    path = config.DATA / f"export-{version}.{kind}"
    if path.exists():
        return path.read_bytes()
    with _build_lock:
        if path.exists():
            return path.read_bytes()
        rows = rows_for(store)
        body = to_xlsx(rows) if kind == "xlsx" else to_csv(rows)
        for old in config.DATA.glob(f"export-*.{kind}"):
            old.unlink(missing_ok=True)
        tmp = path.with_suffix(".tmp")
        tmp.write_bytes(body)
        tmp.replace(path)
        return body
