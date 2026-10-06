"""The web app: accounts, access rules, running, downloads, settings, schedule, headers. No network, no Google."""
import io
import re
import sys
import tempfile
import threading
import time
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from itleads import cli, config, sheet  # noqa: E402
from itleads.store import Store  # noqa: E402
from itleads.web import appdb, auth, create_app, export, jobs  # noqa: E402
from itleads.web.jobs import RunManager, Scheduler  # noqa: E402

PASSWORD = "correct horse battery"
TEAM_PASSWORD = "unit-test-team-pass-2026"          # the program ships none: every test sets its own


class WebCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        root = Path(self._tmp.name)
        self._saved = (config.ROOT, config.DATA, config.LOGS, config.CONFIG_PATH)
        config.ROOT, config.DATA, config.LOGS, config.CONFIG_PATH = root, root / "data", root / "logs", root / "config.json"
        config.ensure_dirs()
        config.update(lambda c: c["app"]["login"].__setitem__("password", TEAM_PASSWORD))
        self.runs = []
        self.gate = threading.Event()
        self.gate.set()
        self.managers = []
        self._fast = mock.patch.object(auth, "METHOD", "pbkdf2:sha256:1000")      # the real cost (600000 rounds) would slow every test
        self._fast.start()
        self.app = self.make_app(self.fake_runner)

    def make_app(self, runner):
        app = create_app(runner=runner)
        app.testing = True
        self.managers.append(app.manager)
        return app

    def tearDown(self):
        self.gate.set()
        for m in self.managers:                          # a run still going would write into the next test's files
            end = time.time() + 10
            while m.snapshot()["running"] and time.time() < end:
                time.sleep(0.02)
        self._fast.stop()
        config.ROOT, config.DATA, config.LOGS, config.CONFIG_PATH = self._saved
        self._tmp.cleanup()

    # a runner that records the call, optionally waits, and writes its run into the database like the real one
    def fake_runner(self, cfg, **kw):
        self.runs.append(kw.get("trigger"))
        kw["progress"]("fetch", 1, 2, "ct")
        self.gate.wait(5)
        store = Store(config.DATA / "leads.db")
        summary = {"status": "ok", "added": 2, "held": 5, "seconds": 3, "sources": {"ct": 2}, "errors": [],
                   "trigger": kw.get("trigger")}
        now = datetime.now().isoformat()
        store.log_run(now, now, summary)
        store.close()
        return summary

    # ---- helpers
    def wait_idle(self, c, timeout=5):
        end = time.time() + timeout
        while time.time() < end:
            if not c.get("/api/status").get_json()["run"]["running"]:
                return
            time.sleep(0.05)
        self.fail("run never finished")

    def token(self, client, path="/login"):
        html = client.get(path).get_data(as_text=True)
        return re.search(r'name="csrf-token" content="([^"]+)"', html).group(1)

    def make_user(self, email="admin@example.com", password=PASSWORD):
        """An account made directly (the app has no sign-up). The first one is the admin."""
        return appdb.create_user(email, "Test", auth.hash_password(password))

    def login(self, client, email, password=PASSWORD, nxt=""):
        t = self.token(client, "/login")
        return client.post("/login", data={"_csrf": t, "email": email, "password": password, "next": nxt})

    def user_client(self, email="admin@example.com", app=None):
        self.make_user(email)
        c = (app or self.app).test_client()
        r = self.login(c, email)
        self.assertEqual(r.status_code, 302, r.get_data(as_text=True)[:300])
        return c

    def api(self, client, path, body=None, method="post"):
        t = re.search(r'name="csrf-token" content="([^"]+)"', client.get("/").get_data(as_text=True)).group(1)
        return getattr(client, method)(path, json=body, headers={"X-CSRF-Token": t})

    def seed(self, name="Acme LLC", site="https://acme.io", registered="2026-10-02"):
        s = Store(config.DATA / "leads.db")
        rec = {"id": "tx:" + name, "source": "tx", "name": name, "trade_name": "", "state": "TX", "city": "Austin",
               "address": "1 Main St", "zip": "78701", "registered": registered, "industry": "Software", "individual": False,
               "emails": [], "people": []}
        enr = {"website": site, "email": "hi@acme.io", "email_from": "company website", "emails": [], "phone": "5122100100",
               "phone_from": "company website", "why": ["full legal name on the site"], "created": "2026-09-30",
               "contact": {"name": "Ada Lovelace", "title": "Founder", "email": "", "linkedin": ""},
               "linkedin": "https://www.linkedin.com/in/ada", "linkedin_company": "", "it": {"level": "strong"}, "domain": "acme.io"}
        s.add(rec, date.today())
        s.record_attempt(rec["id"], enr, "ready", date.today())
        s.close()


class AuthTests(WebCase):
    def test_anonymous_is_sent_to_sign_in(self):
        c = self.app.test_client()
        r = c.get("/")
        self.assertEqual(r.status_code, 200)                              # the public start page, not the dashboard
        self.assertIn("Sign in", r.get_data(as_text=True))
        self.assertNotIn("Run now", r.get_data(as_text=True))
        r = c.get("/settings")
        self.assertEqual((r.status_code, "/login" in r.headers["Location"]), (302, True))
        r = c.get("/api/status")
        self.assertEqual(r.status_code, 401)
        self.assertFalse(r.get_json()["ok"])
        self.assertEqual(c.get("/download.csv").status_code, 302)

    def test_first_account_is_admin_and_later_ones_are_not(self):
        a = self.user_client("one@example.com")
        b = self.user_client("two@example.com")
        users = {u["email"]: u["is_admin"] for u in appdb.list_users()}
        self.assertEqual(users, {"one@example.com": 1, "two@example.com": 0})
        self.assertEqual(a.get("/settings").status_code, 200)
        self.assertEqual(b.get("/settings").status_code, 403)

    def test_passwords_are_hashed_not_stored(self):
        self.user_client()
        h = appdb.user_by_email("admin@example.com")["pw_hash"]
        self.assertNotIn(PASSWORD, h)
        self.assertTrue(h.startswith(("scrypt:", "pbkdf2:")))

    def test_login_and_logout(self):
        self.user_client()
        c = self.app.test_client()
        r = self.login(c, "admin@example.com", "wrong password")
        self.assertEqual(r.status_code, 200)                              # the form again, with its message (not a 401)
        self.assertIn("do not match", r.get_data(as_text=True))
        r = self.login(c, "nobody@example.com")
        self.assertIn("do not match", r.get_data(as_text=True))       # same message: no hint who has an account
        self.assertEqual(self.login(c, "admin@example.com").status_code, 302)
        self.assertEqual(c.get("/").status_code, 200)
        t = re.search(r'name="csrf-token" content="([^"]+)"', c.get("/").get_data(as_text=True)).group(1)
        self.assertEqual(c.post("/logout", data={"_csrf": t}).status_code, 302)
        self.assertEqual(c.get("/settings").status_code, 302)                           # a page that needs an account

    def test_csrf_is_required_on_every_post(self):
        c = self.user_client()
        self.assertEqual(c.post("/logout").status_code, 400)
        self.assertEqual(c.post("/api/run", json={}).status_code, 400)
        self.assertEqual(c.post("/api/settings/schedule", json={"time": "08:00"}).status_code, 400)
        self.assertEqual(c.post("/api/run", json={}, headers={"X-CSRF-Token": "forged"}).status_code, 400)
        anon = self.app.test_client()
        self.assertEqual(anon.post("/login", data={"email": "a@b.co", "password": "x" * 9}).status_code, 400)

    def test_repeated_failures_lock_the_address_out(self):
        self.user_client()
        c = self.app.test_client()
        for _ in range(5):
            self.assertEqual(self.login(c, "admin@example.com", "nope nope nope").status_code, 200)
        r = self.login(c, "admin@example.com")                        # even the right password waits
        self.assertEqual(r.status_code, 429)

    def test_only_local_paths_are_followed_after_sign_in(self):
        self.user_client()
        for nxt, want in (("//evil.example/x", "/"), ("https://evil.example", "/"), ("/\\evil.example", "/"),
                          ("/settings", "/settings")):
            c = self.app.test_client()
            r = self.login(c, "admin@example.com", nxt=nxt)
            self.assertEqual(r.headers["Location"], want, nxt)

    def test_changing_a_password_ends_old_sessions(self):
        c = self.user_client()
        self.assertEqual(c.get("/").status_code, 200)
        u = appdb.user_by_email("admin@example.com")
        appdb.set_password(u["id"], auth.hash_password("a brand new secret"))
        self.assertEqual(c.get("/settings").status_code, 302)

    def test_logging_out_clears_this_browsers_cookie_and_leaves_everyone_else_signed_in(self):
        c = self.user_client()
        other = self.app.test_client()
        self.assertEqual(self.login(other, "admin@example.com").status_code, 302)       # a second person, same account
        t = re.search(r'name="csrf-token" content="([^"]+)"', c.get("/").get_data(as_text=True)).group(1)
        c.post("/logout", data={"_csrf": t})
        self.assertEqual(c.get("/settings").status_code, 302)                           # this browser is signed out
        self.assertEqual(other.get("/settings").status_code, 200)                       # the other person is not
        copied = other.get_cookie("itleads").value
        thief = self.app.test_client()
        thief.set_cookie("itleads", copied)
        appdb.set_password(appdb.user_by_email("admin@example.com")["id"], auth.hash_password("a brand new secret"))
        self.assertEqual(thief.get("/settings").status_code, 302)                       # changing the password ends them all

    def test_resetting_a_password_from_the_command_line_also_lifts_the_lockout(self):
        self.user_client()
        c = self.app.test_client()
        for _ in range(5):
            self.login(c, "admin@example.com", "nope nope nope")
        self.assertEqual(self.login(c, "admin@example.com").status_code, 429)
        appdb.clear_lockout("Admin@Example.com")
        self.assertEqual(self.login(c, "admin@example.com").status_code, 302)

    def test_passwords_use_a_hash_every_python_can_check_and_an_unreadable_one_just_fails(self):
        self._fast.stop()                                                        # look at the real setting, not the test shortcut
        try:
            self.assertEqual(auth.METHOD, "pbkdf2:sha256:600000")
        finally:
            self._fast.start()
        h = auth.hash_password("correct horse battery")
        self.assertTrue(auth._verify(h, "correct horse battery"))
        self.assertFalse(auth._verify(h, "wrong"))
        with mock.patch.object(auth, "check_password_hash", side_effect=AttributeError("no scrypt here")):
            self.assertFalse(auth._verify(h, "correct horse battery"))     # a hash this Python cannot read: no crash



class ThrottleTests(WebCase):
    def post_login(self, email, password, ip="127.0.0.1"):
        c = self.app.test_client()
        t = re.search(r'name="csrf-token" content="([^"]+)"', c.get("/login", environ_overrides={"REMOTE_ADDR": ip}).get_data(as_text=True))
        t = t.group(1) if t else ""
        return c.post("/login", data={"_csrf": t, "email": email, "password": password}, environ_overrides={"REMOTE_ADDR": ip})

    def test_the_throttle_table_stays_small_whatever_is_typed_into_the_email_box(self):
        self.user_client()
        for _ in range(3):
            self.post_login("a" * 50000 + "@example.com", "nope nope nope")
        import sqlite3
        keys = [r[0] for r in sqlite3.connect(config.DATA / "app.db").execute("SELECT k FROM attempts")]
        self.assertTrue(keys)
        self.assertLess(max(len(k) for k in keys), 100)                      # fixed-size stand-ins, never the typed text
        self.assertLess((config.DATA / "app.db").stat().st_size, 200_000)

    def test_parallel_guesses_cannot_slip_under_the_limit(self):
        self.user_client()
        seen = []
        real = auth._verify
        def counting(h, p):
            seen.append(1)
            return real(h, p)
        from concurrent.futures import ThreadPoolExecutor
        with mock.patch.object(auth, "_verify", counting):
            with ThreadPoolExecutor(10) as pool:
                codes = list(pool.map(lambda i: self.post_login("admin@example.com", f"wrong guess {i}").status_code, range(30)))
        self.assertLessEqual(len(seen), auth.LIMIT_PAIR)                      # only five were ever checked
        self.assertGreaterEqual(codes.count(429), 20)

    def test_a_correct_password_does_not_use_up_the_limit(self):
        self.user_client()
        for _ in range(auth.LIMIT_PAIR + 3):
            self.assertEqual(self.post_login("admin@example.com", PASSWORD).status_code, 302)

    def test_an_ipv6_visitor_cannot_dodge_the_limit_by_changing_the_end_of_the_address(self):
        a, b, c = ("2001:db8:1:2:aaaa::1", "2001:db8:1:2:bbbb::9", "2001:db8:1:3::1")
        with self.app.test_request_context(environ_overrides={"REMOTE_ADDR": a}):
            ka = auth.client_key()
        with self.app.test_request_context(environ_overrides={"REMOTE_ADDR": b}):
            kb = auth.client_key()
        with self.app.test_request_context(environ_overrides={"REMOTE_ADDR": c}):
            kc = auth.client_key()
        self.assertEqual(ka, kb)
        self.assertNotEqual(ka, kc)

    def test_an_account_made_by_an_older_version_is_stored_the_current_way_at_the_next_sign_in(self):
        from werkzeug.security import generate_password_hash
        appdb.create_user("old@example.com", "Old", generate_password_hash(PASSWORD, method="pbkdf2:sha256:50"))
        self.assertFalse(appdb.user_by_email("old@example.com")["pw_hash"].startswith(auth.METHOD))
        self.assertEqual(self.post_login("old@example.com", PASSWORD).status_code, 302)
        self.assertTrue(appdb.user_by_email("old@example.com")["pw_hash"].startswith(auth.METHOD))
        self.assertEqual(self.post_login("old@example.com", PASSWORD).status_code, 302)    # and it still works

    def test_when_too_many_password_checks_run_at_once_the_answer_is_a_polite_busy_page(self):
        self.user_client()
        from contextlib import contextmanager
        @contextmanager
        def full():
            raise auth.Busy()
            yield
        with mock.patch.object(auth, "_hashing", full):
            r = self.post_login("admin@example.com", PASSWORD)
        self.assertEqual(r.status_code, 503)
        self.assertIn("try again", r.get_data(as_text=True).lower())


class HostAndConfigTests(WebCase):
    def test_a_request_addressed_to_a_foreign_name_is_refused(self):
        c = self.app.test_client()
        self.assertEqual(c.get("/login", headers={"Host": "attacker.example:8765"}).status_code, 400)
        self.assertEqual(c.get("/login", headers={"Host": "127.0.0.1:8765"}).status_code, 200)
        self.assertEqual(c.get("/login", headers={"Host": "localhost"}).status_code, 200)
        self.assertEqual(c.get("/login", headers={"Host": "[::1]:8765"}).status_code, 200)
        self.assertEqual(c.get("/healthz", headers={"Host": "attacker.example"}).status_code, 400)

    def test_names_can_be_allowed_or_the_check_switched_off_for_a_proxied_server(self):
        config.update(lambda c: c["app"].__setitem__("allowed_hosts", ["leads.example.com"]))
        app = self.make_app(self.fake_runner)
        self.assertEqual(app.test_client().get("/login", headers={"Host": "leads.example.com"}).status_code, 200)
        self.assertEqual(app.test_client().get("/login", headers={"Host": "other.example.com"}).status_code, 400)
        open_app = create_app(runner=self.fake_runner, enforce_hosts=False)
        self.managers.append(open_app.manager)
        self.assertEqual(open_app.test_client().get("/login", headers={"Host": "anything.example"}).status_code, 200)

    def test_the_environment_can_add_the_names_the_app_answers_to(self):
        from itleads.web import extra_hosts
        cfg = {"app": {"allowed_hosts": ["a.example.com"]}}
        with mock.patch.dict("os.environ", {"ITLEADS_ALLOWED_HOSTS": "B.example.com, c.example.com ,"}):
            self.assertEqual(extra_hosts(cfg), {"a.example.com", "b.example.com", "c.example.com"})
            app = self.make_app(self.fake_runner)
            self.assertEqual(app.test_client().get("/login", headers={"Host": "b.example.com"}).status_code, 200)
            self.assertEqual(app.test_client().get("/login", headers={"Host": "evil.example"}).status_code, 400)

    def test_a_damaged_config_does_not_show_the_install_path_to_visitors(self):
        self.user_client()
        config.CONFIG_PATH.write_text("{ not json")
        r = self.app.test_client().get("/login")                              # anyone can open this page
        self.assertEqual(r.status_code, 500)
        self.assertNotIn(str(config.ROOT), r.get_data(as_text=True))
        self.assertIn("cannot be read", r.get_data(as_text=True))

    def test_a_config_written_loosely_by_an_older_version_is_tightened_on_start(self):
        import os
        config.save({**config.load()})
        os.chmod(config.CONFIG_PATH, 0o644)
        self.make_app(self.fake_runner)
        self.assertEqual(config.CONFIG_PATH.stat().st_mode & 0o777, 0o600)


class LandingAndLoginTests(WebCase):
    def test_visitors_see_the_start_page_and_signed_in_people_the_dashboard(self):
        page = self.app.test_client().get("/").get_data(as_text=True)
        self.assertIn("Hybrid Leads", page)
        self.assertIn("Sign in", page)
        self.assertIn("Internal use only", page)
        self.assertIn('rel="icon" type="image/svg+xml"', page)
        self.assertNotIn("Run now", page)
        c = self.user_client()
        home = c.get("/").get_data(as_text=True)
        self.assertIn("New IT company filings", home)
        self.assertIn("Run now", home)

    def test_there_is_no_sign_up(self):
        c = self.app.test_client()
        r = c.get("/signup")
        self.assertEqual((r.status_code, r.headers["Location"]), (302, "/login"))
        self.assertIn(c.post("/signup", data={"email": "a@b.co", "password": PASSWORD + "x"}).status_code, (400, 405))
        self.assertEqual(appdb.count_users(), 0)                                  # nothing was created
        self.assertNotIn("Create an account", c.get("/login").get_data(as_text=True))
        self.assertNotIn("Create an account", c.get("/").get_data(as_text=True))

    def test_the_shared_login_is_not_written_on_any_page_by_default_and_it_works(self):
        cfg = config.load()
        auth.ensure_internal_login(cfg)
        lg = auth.internal_login(cfg)
        self.assertFalse(lg["show"])
        self.assertGreaterEqual(len(lg["password"]), 16)
        for path in ("/", "/login"):
            page = self.app.test_client().get(path).get_data(as_text=True)
            self.assertNotIn(lg["password"], page, path)
            self.assertNotIn("hint-password", page, path)
        c = self.app.test_client()
        self.assertEqual(self.login(c, lg["email"], "not-the-team-password-0").status_code, 200)  # a wrong one is refused: the form again
        r = self.login(c, lg["email"], lg["password"])
        self.assertEqual(r.status_code, 302)
        self.assertEqual(c.get("/settings").status_code, 200)                    # the shared account is an admin

    def test_the_program_ships_without_a_team_password(self):
        self.assertEqual(config.DEFAULTS["app"]["login"]["password"], "")
        source = (ROOT / "itleads" / "config.py").read_text(encoding="utf-8")
        self.assertNotRegex(source, r'"password":\s*"[^"]')                         # no password is written in the code

    def test_without_a_password_nobody_can_sign_in_and_the_start_up_line_says_how_to_fix_it(self):
        config.update(lambda c: c["app"]["login"].__setitem__("password", ""))
        auth.ensure_internal_login(config.load())
        self.assertEqual(appdb.count_users(), 0)                                   # no account is made
        self.assertEqual(self.login(self.app.test_client(), "team@hybrid.agency", "anything-at-all").status_code, 200)
        self.assertIn("set-login", cli.login_line(auth.internal_login(config.load())))

    def test_the_start_up_line_never_prints_the_password(self):
        lg = auth.internal_login(config.load())
        line = cli.login_line(lg)
        self.assertIn(lg["email"], line)
        self.assertNotIn(lg["password"], line)
        self.assertNotIn(lg["password"], cli.login_line({**lg, "show": True}))

    def test_the_login_can_be_written_on_the_page_when_asked_for(self):
        with mock.patch.dict("os.environ", {"ITLEADS_SHOW_LOGIN": "1"}):
            lg = auth.internal_login(config.load())
            self.assertTrue(lg["show"])
            self.assertIn(lg["password"], self.app.test_client().get("/login").get_data(as_text=True))
        self.assertFalse(auth.internal_login(config.load())["show"])

    def test_the_login_can_be_changed_or_hidden(self):
        config.update(lambda c: c["app"].__setitem__("login", {"email": "crew@agency.test", "password": "Another-Pass-1", "show": True}))
        page = self.app.test_client().get("/login").get_data(as_text=True)
        self.assertIn("crew@agency.test", page)
        self.assertIn("Another-Pass-1", page)
        config.update(lambda c: c["app"]["login"].__setitem__("show", False))
        self.assertNotIn("Another-Pass-1", self.app.test_client().get("/login").get_data(as_text=True))
        config.update(lambda c: c["app"]["login"].__setitem__("show", True))
        with mock.patch.dict("os.environ", {"ITLEADS_SHOW_LOGIN": "0"}):
            self.assertNotIn("Another-Pass-1", self.app.test_client().get("/login").get_data(as_text=True))
        with mock.patch.dict("os.environ", {"ITLEADS_LOGIN_EMAIL": "ops@agency.test", "ITLEADS_LOGIN_PASSWORD": "From-Env-Pass-9"}):
            page = self.app.test_client().get("/login").get_data(as_text=True)
            self.assertIn("ops@agency.test", page)
            self.assertIn("From-Env-Pass-9", page)

    def test_the_shared_account_is_made_an_admin_and_given_the_written_password_at_every_start(self):
        cfg = config.load()
        self.make_user("someone@example.com")                                    # an older account: the first one, an admin
        auth.ensure_internal_login(cfg)
        user = appdb.user_by_email("team@hybrid.agency")
        self.assertEqual(user["is_admin"], 1)
        again = appdb.user_by_email("team@hybrid.agency")
        auth.ensure_internal_login(cfg)                                          # idempotent
        self.assertEqual(appdb.user_by_email("team@hybrid.agency")["pw_hash"], again["pw_hash"])
        appdb.set_password(user["id"], auth.hash_password("something else"))     # somebody changed it: the written one wins
        appdb.touch_login(user["id"])
        auth.ensure_internal_login(cfg)
        c = self.app.test_client()
        self.assertEqual(self.login(c, "team@hybrid.agency", cfg["app"]["login"]["password"]).status_code, 302)
        appdb.set_password(user["id"], auth.hash_password("x" * 12))
        appdb.make_admin(user["id"])
        self.assertEqual(appdb.user_by_email("someone@example.com")["is_admin"], 1)

    def test_changing_the_shared_email_retires_the_old_login_but_not_other_accounts(self):
        cfg = config.load()
        self.make_user("someone@example.com")                                    # somebody's own account: never touched
        auth.ensure_internal_login(cfg)
        self.assertIsNotNone(appdb.user_by_email("team@hybrid.agency"))
        config.update(lambda c: c["app"]["login"].__setitem__("email", "agency@hybrid.agency"))
        auth.ensure_internal_login(config.load())
        self.assertIsNone(appdb.user_by_email("team@hybrid.agency"))             # the old written login no longer works
        self.assertIsNotNone(appdb.user_by_email("agency@hybrid.agency"))
        self.assertIsNotNone(appdb.user_by_email("someone@example.com"))
        c = self.app.test_client()
        self.assertEqual(self.login(c, "team@hybrid.agency", cfg["app"]["login"]["password"]).status_code, 200)   # refused: the form again
        self.assertEqual(self.login(c, "agency@hybrid.agency", cfg["app"]["login"]["password"]).status_code, 302)

    def test_starting_the_server_seeds_the_login(self):
        create_app(runner=self.fake_runner, seed_login=True)
        self.assertIsNotNone(appdb.user_by_email("team@hybrid.agency"))
        self.assertEqual(appdb.count_users(), 1)

    def test_the_hybrid_icon_is_served_and_linked(self):
        c = self.app.test_client()
        for path, kind in (("/static/favicon.svg", "image/svg+xml"), ("/static/mark.svg", "image/svg+xml"),
                           ("/static/favicon-32.png", "image/png"), ("/static/favicon-48.png", "image/png"),
                           ("/static/apple-touch-icon.png", "image/png"), ("/static/icon-192.png", "image/png"),
                           ("/static/icon-512.png", "image/png"), ("/static/hybrid-logo.png", "image/png")):
            r = c.get(path)
            self.assertEqual(r.status_code, 200, path)
            self.assertEqual(r.mimetype, kind, path)
        self.assertIn(b'fill="#7BDD4C"', c.get("/static/favicon.svg").data)       # the Hybrid lime
        for svg in ("/static/favicon.svg", "/static/mark.svg"):                    # the content policy forbids inline styles and scripts
            body = c.get(svg).get_data(as_text=True).lower()
            for bad in ("<style", "<script", " style=", " onload=", "<foreignobject", "href="):
                self.assertNotIn(bad, body, svg + " " + bad)
        head = c.get("/login").get_data(as_text=True)
        for needle in ("favicon.svg", "favicon-32.png", "favicon-48.png", "apple-touch-icon.png", "site.webmanifest", "mark.svg"):
            self.assertIn(needle, head)
        manifest = c.get("/static/site.webmanifest")
        self.assertEqual(manifest.status_code, 200)
        for icon in manifest.get_json()["icons"]:                                  # every icon the manifest names exists
            self.assertEqual(c.get(icon["src"]).status_code, 200, icon["src"])
        r = c.get("/favicon.ico")                                                  # asked for by habit, signed in or not
        self.assertEqual((r.status_code, r.mimetype), (200, "image/png"))
        self.assertTrue(r.data.startswith(b"\x89PNG"))

    def test_settings_has_no_people_list_and_no_sign_up_switch(self):
        admin = self.user_client()
        html = admin.get("/settings").get_data(as_text=True)
        self.assertNotIn("People", html)
        self.assertNotIn("create an account", html.lower())
        self.assertEqual(self.api(admin, "/api/settings/access", {"signup_open": True}).status_code, 404)
        self.assertEqual(admin.get("/account").status_code, 404)


class SecurityTests(WebCase):
    def test_headers_and_cookie(self):
        c = self.app.test_client()
        r = c.get("/login")
        csp = r.headers["Content-Security-Policy"]
        for part in ("default-src 'self'", "frame-src 'none'", "frame-ancestors 'none'", "script-src 'self'"):
            self.assertIn(part, csp)
        self.assertEqual(r.headers["X-Frame-Options"], "DENY")
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        self.make_user()
        r = self.login(c, "admin@example.com")
        cookie = r.headers.get("Set-Cookie", "")
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=Lax", cookie)
        self.assertNotIn("Secure", cookie)                                               # plain http on this Mac
        self.assertEqual(c.get("/").headers["Cache-Control"], "no-store")
        self.assertEqual(c.get("/settings").headers["Cache-Control"], "no-store")

    def test_https_mode_turns_on_secure_cookies_and_hsts(self):
        with mock.patch.dict("os.environ", {"ITLEADS_SECURE_COOKIES": "1"}):
            app = create_app(runner=self.fake_runner)
            app.testing = True
            c = app.test_client()
            self.make_user("a@example.com")
            r = c.get("/login")
            self.assertIn("max-age=", r.headers["Strict-Transport-Security"])
            t = re.search(r'name="csrf-token" content="([^"]+)"', r.get_data(as_text=True)).group(1)
            r = c.post("/login", data={"_csrf": t, "email": "a@example.com", "password": PASSWORD})
            self.assertIn("Secure", r.headers.get("Set-Cookie", ""))

    def test_no_page_uses_inline_styles_or_scripts(self):
        c = self.user_client()
        self.seed()
        for path in ("/", "/settings", "/fragment/stats"):
            html = c.get(path).get_data(as_text=True)
            self.assertNotRegex(html, r'\sstyle="', path)
            self.assertNotRegex(html, r"<script(?![^>]*\bsrc=)", path)
        anon = self.app.test_client()
        for path in ("/", "/login"):
            html = anon.get(path).get_data(as_text=True)
            self.assertNotRegex(html, r'\sstyle="', path)
            self.assertNotRegex(html, r"<script(?![^>]*\bsrc=)", path)

    def test_untrusted_text_is_escaped(self):
        c = self.user_client()
        self.seed(name='<img src=x onerror=alert(1)> LLC')
        r = c.get("/api/preview")                                     # JSON is data: the page inserts it as text only
        self.assertEqual(r.mimetype, "application/json")
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        js = (ROOT / "itleads/web/static/app.js").read_text()
        self.assertEqual(js.count("innerHTML"), 1)                      # only our own server-rendered fragment
        self.assertIn("textContent = v", js)                            # table cells get text, never markup
        for path in ("/", "/fragment/stats", "/settings"):
            self.assertNotIn("<img src=x", c.get(path).get_data(as_text=True), path)

    def test_errors_are_json_for_the_api_and_html_for_pages(self):
        c = self.user_client()
        self.assertEqual(c.get("/api/nothing").get_json()["ok"], False)
        r = c.get("/nothing-here")
        self.assertEqual(r.status_code, 404)
        self.assertIn("Page not found", r.get_data(as_text=True))

    def test_settings_apis_are_admin_only(self):
        self.user_client("admin@example.com")
        member = self.user_client("member@example.com")
        for path, body in (("/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"}),
                           ("/api/settings/sharing", {"link": False}), ("/api/settings/schedule", {"time": "06:00"}),
                           ("/api/settings/rules", {"require": ["email"], "sources": {"tx": True}})):
            r = self.api(member, path, body)
            self.assertEqual(r.status_code, 403, path)


class RunTests(WebCase):
    def test_run_starts_once_and_reports(self):
        c = self.user_client()
        self.gate.clear()
        r = self.api(c, "/api/run", {})
        self.assertEqual(r.status_code, 202)
        r2 = self.api(c, "/api/run", {})
        self.assertEqual(r2.status_code, 409)
        snap = c.get("/api/status").get_json()["run"]
        self.assertTrue(snap["running"])
        self.assertEqual((snap["stage"], snap["done"], snap["total"]), ("fetch", 1, 2))
        self.gate.set()
        self.wait_idle(c)
        last = c.get("/api/status").get_json()["run"]["last"]
        self.assertEqual(last["status"], "ok")
        self.assertEqual(self.runs, ["manual"])

    def test_members_wait_between_manual_runs_but_admins_do_not(self):
        admin = self.user_client("admin@example.com")
        member = self.user_client("member@example.com")
        self.assertEqual(self.api(member, "/api/run", {}).status_code, 202)
        self.wait_idle(member)
        r = self.api(member, "/api/run", {})
        self.assertEqual(r.status_code, 429)
        self.assertIn("can run it again", r.get_json()["message"])
        self.assertEqual(self.api(admin, "/api/run", {}).status_code, 202)
        self.wait_idle(admin)

    def test_a_crashed_run_still_starts_the_cooldown(self):
        def boom(cfg, **kw):
            raise RuntimeError("registry exploded")
        app = self.make_app(boom)
        self.user_client("admin@example.com", app)
        member = self.user_client("member@example.com", app)
        t = re.search(r'name="csrf-token" content="([^"]+)"', member.get("/").get_data(as_text=True)).group(1)
        self.assertEqual(member.post("/api/run", json={}, headers={"X-CSRF-Token": t}).status_code, 202)
        for _ in range(60):
            if not member.get("/api/status").get_json()["run"]["running"]:
                break
            time.sleep(0.05)
        r = member.post("/api/run", json={}, headers={"X-CSRF-Token": t})
        self.assertEqual(r.status_code, 429)
        self.assertIn("can run it again", r.get_json()["message"])

    def test_a_crashing_run_is_reported_not_fatal(self):
        def boom(cfg, **kw):
            raise RuntimeError("registry exploded")
        app = self.make_app(boom)
        c = self.user_client(app=app)
        t = re.search(r'name="csrf-token" content="([^"]+)"', c.get("/").get_data(as_text=True)).group(1)
        self.assertEqual(c.post("/api/run", json={}, headers={"X-CSRF-Token": t}).status_code, 202)
        for _ in range(60):
            snap = c.get("/api/status").get_json()["run"]
            if not snap["running"]:
                break
            time.sleep(0.05)
        self.assertIn("registry exploded", snap["last"]["error"])
        self.assertEqual(c.get("/").status_code, 200)                 # the app is fine


class DataTests(WebCase):
    def test_dashboard_shows_what_is_there(self):
        c = self.user_client()
        html = c.get("/").get_data(as_text=True)
        self.assertIn("No companies yet", html)
        self.seed()
        html = c.get("/").get_data(as_text=True)
        self.assertIn("Total listed", html)
        self.assertNotIn("No companies yet", html)
        self.assertIn("Texas", html)

    def log_run(self, status, hours_ago, **extra):
        store = Store(config.DATA / "leads.db")
        start = datetime.now() - timedelta(hours=hours_ago)
        store.log_run(start.isoformat(), (start + timedelta(minutes=2)).isoformat(),
                      dict({"status": status, "errors": [], "trigger": "schedule"}, **extra))
        store.close()

    def test_updated_means_the_last_run_that_really_refreshed_the_list(self):
        c = self.user_client()
        self.log_run("ok", 30, added=3, pushed=3, held=9, seconds=60, sources={"ct": 3})
        self.log_run("offline", 2, errors=["no internet connection"])
        page = c.get("/").get_data(as_text=True)
        good = (datetime.now() - timedelta(hours=30) + timedelta(minutes=2)).strftime("%H:%M")      # the run FINISHED two minutes after it began
        self.assertIn(f"{good}", page)                                              # 'Updated' shows the ok run's time
        self.assertIn("Last run: no internet", page)                                # and the failed attempt is said apart
        self.assertIn("Nothing changed", page)                                      # the history row has no made-up counts
        self.assertNotIn("0 added", page.split("Recent runs")[1].split("Nothing changed")[0].split("</tr>")[0])
        status = c.get("/api/status").get_json()
        self.assertEqual(status["problem"], "Last run: no internet")
        self.assertTrue(status["updated"])

    def test_an_empty_list_after_runs_says_so_instead_of_asking_for_a_first_run(self):
        c = self.user_client()
        self.log_run("ok", 3, added=0, pushed=0, held=40, seconds=60)
        page = c.get("/").get_data(as_text=True)
        self.assertIn("Nothing is listed yet", page)
        self.assertNotIn("The first run covers", page)

    def test_csv_download(self):
        c = self.user_client()
        self.seed()
        self.seed(name="=SUM(1,1) Consulting LLC", site="https://sum.io")
        r = c.get("/download.csv")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment", r.headers["Content-Disposition"])
        text = r.get_data(as_text=True)
        self.assertTrue(text.startswith("﻿Company,Website,Email,Phone,Address,Registered"))
        self.assertIn("Acme LLC", text)
        self.assertIn("'=SUM(1,1) Consulting LLC", text)                # a spreadsheet would run this otherwise

    def test_xlsx_download(self):
        c = self.user_client()
        self.seed()
        self.seed(name="=SUM(1,1) Consulting LLC", site="https://sum.io")
        r = c.get("/download.xlsx")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.get_data())).active
        self.assertEqual([c_.value for c_ in ws[1]][:6], ["COMPANY", "WEBSITE", "EMAIL", "PHONE", "ADDRESS", "REGISTERED"])
        rows = {ws.cell(row=i, column=1).value: i for i in range(2, ws.max_row + 1)}
        self.assertIn("Acme LLC", rows)
        i = rows["Acme LLC"]
        self.assertEqual(ws.cell(row=i, column=2).value, "acme.io")
        self.assertEqual(ws.cell(row=i, column=2).hyperlink.target, "https://acme.io")
        self.assertEqual(ws.cell(row=i, column=3).hyperlink.target, "mailto:hi@acme.io")
        self.assertEqual(ws.cell(row=i, column=4).value, "(512) 210-0100")
        self.assertEqual(ws.cell(row=i, column=6).value.date(), date(2026, 10, 2))
        evil = ws.cell(row=rows["=SUM(1,1) Consulting LLC"], column=1)
        self.assertEqual(evil.data_type, "s")                           # stored as text, never a formula
        self.assertEqual(ws.freeze_panes, "B2")

    def test_downloads_are_built_once_per_change_of_data(self):
        c = self.user_client()
        self.seed()
        with mock.patch.object(export, "to_xlsx", wraps=export.to_xlsx) as spy:
            first = c.get("/download.xlsx").get_data()
            second = c.get("/download.xlsx").get_data()
            self.assertEqual((spy.call_count, first == second), (1, True))
            self.seed(name="Beta LLC", site="https://beta.io")
            self.assertGreater(len(c.get("/download.xlsx").get_data()), 0)
            self.assertEqual(spy.call_count, 2)
        self.assertEqual(len(list(config.DATA.glob("export-*.xlsx"))), 1)                # old builds are cleaned up

    def test_downloads_work_with_no_data_and_reject_other_formats(self):
        c = self.user_client()
        self.assertEqual(c.get("/download.xlsx").status_code, 200)
        self.assertEqual(c.get("/download.csv").status_code, 200)
        self.assertEqual(c.get("/download.exe").status_code, 404)

    def test_preview_api_before_and_after_google(self):
        c = self.user_client()
        self.seed()
        d = c.get("/api/preview").get_json()
        self.assertFalse(d["connected"])
        self.assertEqual(d["total"], 1)
        self.assertEqual(d["rows"][0]["company"], "Acme LLC")
        config.update(lambda cfg: cfg.update(apps_script_url="https://script.google.com/macros/s/x/exec",
                                             sheet_id="SHEET123", sheet_url="javascript:alert(1)"))
        d = c.get("/api/preview").get_json()
        self.assertTrue(d["connected"] and d["linked"])
        self.assertEqual(d["embed_url"], "https://docs.google.com/spreadsheets/d/SHEET123/htmlview")
        self.assertEqual(d["sheet_url"], "https://docs.google.com/spreadsheets/d/SHEET123/edit")   # built, never trusted
        config.update(lambda cfg: cfg.update(share_link=False))
        self.assertFalse(c.get("/api/preview").get_json()["linked"])
        config.update(lambda cfg: cfg.update(sheet_id="a b<script>"))
        self.assertFalse(c.get("/api/preview").get_json()["connected"])                       # a odd id is never embedded

    def test_preview_is_capped_and_says_how_many_there_are(self):
        c = self.user_client()
        for i in range(3):
            self.seed(name=f"Co {i} LLC", site=f"https://co{i}.io", registered=f"2026-10-0{i + 1}")
        with mock.patch("itleads.web.views.PREVIEW_ROWS", 2):
            d = c.get("/api/preview").get_json()
        self.assertEqual((d["total"], len(d["rows"])), (3, 2))
        self.assertEqual(d["rows"][0]["company"], "Co 2 LLC")                                  # newest first


class FakeBridge:
    """Stands in for the Google script."""
    calls = []
    fail = None

    id = "SHEET-ID-42"
    refuse_link = False                      # an account whose rules forbid sharing by link

    def __init__(self, url, token, timeout=0, retries=3):
        self.url, self.token, self.retries = url, token, retries

    def ping(self):
        if FakeBridge.fail:
            raise sheet.BridgeError(FakeBridge.fail)
        return {"ok": True}

    def init(self, tz, share, refresh=False, link=None):
        FakeBridge.calls.append(("init", link))
        refused = bool(link) and FakeBridge.refuse_link
        return {"ok": True, "id": FakeBridge.id, "url": "https://docs.google.com/spreadsheets/d/x/edit",
                "shared": bool(link) and not refused, "share_error": "Access denied: DriveApp." if refused else ""}

    def share(self, link):
        FakeBridge.calls.append(("share", link))
        refused = bool(link) and FakeBridge.refuse_link
        return {"ok": True, "shared": bool(link) and not refused, "share_error": "Access denied: DriveApp." if refused else ""}


class SettingsTests(WebCase):
    def setUp(self):
        super().setUp()
        FakeBridge.calls, FakeBridge.fail, FakeBridge.id, FakeBridge.refuse_link = [], None, "SHEET-ID-42", False
        self.sent = []
        self.patch = mock.patch.object(sheet, "Bridge", FakeBridge)
        self.patch.start()
        self.admin = self.user_client()

    def tearDown(self):
        self.patch.stop()
        super().tearDown()

    def test_connect_google(self):
        good = "https://script.google.com/macros/s/AKfycbx_abc-123/exec"
        r = self.api(self.admin, "/api/settings/google", {"url": good})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        cfg = config.load()
        self.assertEqual((cfg["sheet_id"], cfg["apps_script_url"]), ("SHEET-ID-42", good))
        self.assertEqual(cfg["sheet_url"], "https://docs.google.com/spreadsheets/d/SHEET-ID-42/edit")
        self.assertTrue(cfg["token"])
        self.assertEqual(FakeBridge.calls[0], ("init", True))          # link sharing on by default for the preview
        self.assertFalse(r.get_json()["sending"])                      # nothing was waiting

    def test_connecting_with_link_sharing_off_keeps_the_sheet_private(self):
        r = self.api(self.admin, "/api/settings/google",
                     {"url": "https://script.google.com/macros/s/x/exec", "link": False})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(FakeBridge.calls[0], ("init", False))
        self.assertFalse(config.load()["share_link"])

    def test_connecting_sends_the_companies_that_were_waiting(self):
        self.seed()
        def sender(cfg, **kw):
            self.sent.append((kw["trigger"], cfg["sheet_id"]))
            return {"status": "ok", "pushed": 1, "added": 0, "held": 0, "errors": [], "trigger": kw["trigger"]}
        with mock.patch.object(self.app.manager, "_sender", sender):
            r = self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
            self.assertEqual((r.status_code, r.get_json()["waiting"], r.get_json()["sending"]), (200, 1, True))
            self.wait_idle(self.admin)
        self.assertEqual(self.sent, [("connect", "SHEET-ID-42")])

    def test_connecting_during_a_run_sends_right_after_it_ends(self):
        self.seed()
        calls = []
        def sender(cfg, **kw):
            calls.append(kw["trigger"])
            return {"status": "ok", "pushed": 1, "added": 0, "held": 0, "errors": [], "trigger": kw["trigger"]}
        self.gate.clear()
        self.assertEqual(self.api(self.admin, "/api/run", {}).status_code, 202)
        with mock.patch.object(self.app.manager, "_sender", sender):
            r = self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
            self.assertEqual((r.get_json()["sending"], r.get_json()["queued"]), (True, True))
            self.assertEqual(calls, [])                                    # not while the run is going
            self.gate.set()
            for _ in range(200):
                if calls:
                    break
                time.sleep(0.02)
            self.wait_idle(self.admin)
        self.assertEqual(calls, ["connect"])

    def test_connecting_still_works_when_the_account_forbids_link_sharing(self):
        FakeBridge.refuse_link = True
        r = self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
        self.assertEqual(r.status_code, 200)
        body = r.get_json()
        self.assertIn("plain table", body["warning"])
        self.assertIn("Access denied", body["warning"])
        self.assertFalse(body["google"]["linked"])                      # so the preview will not try to embed it
        cfg = config.load()
        self.assertEqual(cfg["sheet_id"], "SHEET-ID-42")
        self.assertFalse(cfg["share_link"])
        d = self.admin.get("/api/preview").get_json()
        self.assertTrue(d["connected"] and not d["linked"])

    def test_turning_link_sharing_on_reports_a_refusal_and_changes_nothing(self):
        self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec", "link": False})
        FakeBridge.refuse_link = True
        r = self.api(self.admin, "/api/settings/sharing", {"link": True})
        self.assertEqual(r.status_code, 400)
        self.assertIn("would not let", r.get_json()["error"])
        self.assertFalse(config.load()["share_link"])
        FakeBridge.refuse_link = False
        self.assertEqual(self.api(self.admin, "/api/settings/sharing", {"link": True}).status_code, 200)
        self.assertTrue(config.load()["share_link"])

    def test_a_different_sheet_gets_everything_again_but_the_same_sheet_does_not(self):
        self.seed()
        store = Store(config.DATA / "leads.db")
        store.mark_pushed("tx:Acme LLC", "hash", date.today())               # it was sent to an earlier sheet
        store.close()
        calls = []
        def sender(cfg, **kw):
            calls.append(cfg["sheet_id"])
            return {"status": "ok", "pushed": 1, "added": 0, "held": 0, "errors": [], "trigger": kw["trigger"]}
        url = {"url": "https://script.google.com/macros/s/x/exec"}
        with mock.patch.object(self.app.manager, "_sender", sender):
            config.update(lambda c: c.update(sheet_id="OLD-SHEET-1", apps_script_url=url["url"]))
            r = self.api(self.admin, "/api/settings/google", url)                  # a new sheet id: start over
            self.assertEqual((r.get_json()["waiting"], r.get_json()["sending"]), (1, True))
            self.wait_idle(self.admin)
            self.assertEqual(calls, ["SHEET-ID-42"])
            store = Store(config.DATA / "leads.db")
            store.mark_pushed("tx:Acme LLC", "hash", date.today())               # now it is in the new sheet
            store.close()
            r = self.api(self.admin, "/api/settings/google", url)                  # same sheet again (script updated)
            self.assertEqual((r.get_json()["waiting"], r.get_json()["sending"]), (0, False))
        self.assertEqual(calls, ["SHEET-ID-42"])

    def test_an_odd_sheet_id_from_google_is_refused(self):
        FakeBridge.id = "../../evil\"><script>"
        r = self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(config.load()["sheet_id"], "")

    def test_settings_reject_malformed_payloads(self):
        for path, body in (("/api/settings/rules", {"require": "email", "sources": {"tx": True}}),
                           ("/api/settings/rules", {"require": ["email"], "sources": []}),
                           ("/api/settings/schedule", {"time": 7}), ("/api/settings/schedule", {"time": ["07:00"]}),
                           ("/api/settings/google", {"url": ["https://script.google.com/macros/s/x/exec"]})):
            self.assertEqual(self.api(self.admin, path, body).status_code, 400, (path, body))
        t = re.search(r'name="csrf-token" content="([^"]+)"', self.admin.get("/").get_data(as_text=True)).group(1)
        r = self.admin.post("/api/settings/rules", data="[1, 2]", content_type="application/json", headers={"X-CSRF-Token": t})
        self.assertEqual(r.status_code, 400)
        self.assertFalse(r.get_json()["ok"])

    def test_only_real_apps_script_addresses_are_accepted(self):
        for bad in ("http://169.254.169.254/latest/meta-data", "https://evil.example/macros/s/x/exec",
                    "https://script.google.com.evil.example/macros/s/x/exec", "file:///etc/passwd", "", "javascript:1"):
            r = self.api(self.admin, "/api/settings/google", {"url": bad})
            self.assertEqual(r.status_code, 400, bad)
        self.assertEqual(config.load()["apps_script_url"], "")

    def test_a_google_error_is_shown_plainly(self):
        FakeBridge.fail = "Google answered with a web page"
        r = self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("web page", r.get_json()["error"])
        self.assertEqual(config.load()["sheet_id"], "")

    def test_sharing_toggle(self):
        self.api(self.admin, "/api/settings/google", {"url": "https://script.google.com/macros/s/x/exec"})
        r = self.api(self.admin, "/api/settings/sharing", {"link": False})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(config.load()["share_link"])
        self.assertIn(("share", False), FakeBridge.calls)

    def test_schedule(self):
        self.assertEqual(self.api(self.admin, "/api/settings/schedule", {"time": "25:00"}).status_code, 400)
        self.assertEqual(self.api(self.admin, "/api/settings/schedule", {"time": "banana"}).status_code, 400)
        r = self.api(self.admin, "/api/settings/schedule", {"time": "06:45"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(config.load()["schedule"], {"hour": 6, "minute": 45})

    def test_rules(self):
        r = self.api(self.admin, "/api/settings/rules",
                     {"require": ["email"], "require_it_signal": False, "sources": {"ct": True, "bogus": True}})
        self.assertEqual(r.status_code, 200)
        cfg = config.load()
        self.assertEqual(cfg["require"], ["website", "email"])         # a website is always required
        self.assertEqual(cfg["sources"], {"tx": False, "ct": True, "seattle": False, "sf": False, "la": False})
        self.assertFalse(cfg["require_it_signal"])
        self.assertEqual(self.api(self.admin, "/api/settings/rules", {"require": [], "sources": {}}).status_code, 400)

    def test_changing_the_rules_releases_companies_that_were_held_back(self):
        store = Store(config.DATA / "leads.db")
        rec = {"id": "tx:H", "source": "tx", "name": "Held LLC", "state": "TX", "city": "Austin", "address": "1 Main St",
               "zip": "78701", "registered": "2026-10-02", "industry": "Software", "individual": False, "emails": [], "people": []}
        store.add(rec, date.today())
        store.record_attempt("tx:H", {"website": "https://held.io", "email": "hi@held.io", "phone": "", "domain": "held.io",
                                      "it": {"level": "strong", "score": 9}, "missing": ["phone"]}, "held", date.today())
        store.close()
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 0)
        r = self.api(self.admin, "/api/settings/rules",
                     {"require": ["email"], "require_it_signal": True, "sources": {"tx": True}})
        self.assertEqual((r.get_json()["released"], r.get_json()["sending"]), (1, False))   # Google is not connected here
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 1)
        again = self.api(self.admin, "/api/settings/rules",
                         {"require": ["email"], "require_it_signal": True, "sources": {"tx": True}})
        self.assertEqual(again.get_json()["released"], 0)

    def test_the_either_contact_rule_releases_companies_that_publish_only_an_email(self):
        store = Store(config.DATA / "leads.db")
        rec = {"id": "ct:K", "source": "ct", "name": "Tarvane LLC", "state": "CT", "city": "Greenwich", "address": "1 Main St",
               "zip": "06830", "registered": "2026-09-22", "industry": "Software", "individual": False, "emails": [], "people": []}
        store.add(rec, date.today())
        store.record_attempt("ct:K", {"website": "https://tarvane.io", "email": "hello@tarvane.io", "phone": "", "domain": "tarvane.io",
                                      "it": {"level": "strong", "score": 9}, "missing": ["phone"], "why": [], "contact": {}}, "held", date.today())
        store.close()
        body = {"require": ["email", "phone"], "require_it_signal": True, "sources": {"ct": True}}
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 0)
        r = self.api(self.admin, "/api/settings/rules", dict(body, contact_either=True))
        self.assertEqual(r.get_json()["released"], 1)
        self.assertTrue(config.load()["contact_either"])
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 1)
        self.assertIn('id="r-either" checked', self.admin.get("/settings").get_data(as_text=True))

    def test_the_skip_established_rule_is_saved_and_applied_at_once(self):
        store = Store(config.DATA / "leads.db")
        rec = {"id": "tx:E", "source": "tx", "name": "Old Co LLC", "state": "TX", "city": "Austin", "address": "1 Main St",
               "zip": "78701", "registered": "2026-09-22", "industry": "Software", "individual": False, "emails": [], "people": []}
        store.add(rec, date.today())
        store.record_attempt("tx:E", {"website": "https://oldco.io", "email": "hi@oldco.io", "phone": "5124460100", "domain": "oldco.io",
                                      "created": "2001-08-05", "it": {"level": "strong", "score": 9}, "why": [], "contact": {}}, "ready", date.today())
        store.close()
        body = {"require": ["email", "phone"], "require_it_signal": True, "sources": {"tx": True}}
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 1)
        self.assertEqual(self.api(self.admin, "/api/settings/rules", dict(body, skip_established=True)).status_code, 200)
        self.assertTrue(config.load()["skip_established"])
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 0)           # gone from the list straight away
        self.assertIn('id="r-old" checked', self.admin.get("/settings").get_data(as_text=True))
        self.api(self.admin, "/api/settings/rules", dict(body, skip_established=False))
        self.assertFalse(config.load()["skip_established"])
        self.assertEqual(self.admin.get("/api/preview").get_json()["total"], 1)           # and back when switched off

    def test_settings_page_shows_the_script_with_its_secret_only_to_admins(self):
        html = self.admin.get("/settings").get_data(as_text=True)
        token = config.load()["token"]
        self.assertTrue(token)
        self.assertIn(token, html)
        member = self.user_client("member@example.com")
        self.assertNotIn(token, member.get("/").get_data(as_text=True))


class SchedulerTests(WebCase):
    # far-future dates: whatever the real clock says, no run on record is "after" the daily time
    D5, D6 = (2099, 1, 5), (2099, 1, 6)

    def at(self, day, h, m=0):
        return datetime(*day, h, m)

    def connect_google(self):
        config.update(lambda c: c.update(apps_script_url="https://script.google.com/macros/s/x/exec", sheet_id="SHEET-ID-42"))

    def wait_manager(self):
        for _ in range(100):
            if not self.app.manager.snapshot()["running"]:
                return
            time.sleep(0.05)
        self.fail("run never finished")

    def test_waits_for_an_account_and_for_a_first_run_or_google(self):
        sched = Scheduler(self.app.manager)
        self.assertEqual(sched.tick(self.at(self.D5, 8)), "not ready")                  # nobody has an account
        self.user_client()
        self.assertEqual(sched.tick(self.at(self.D5, 8)), "not ready")                  # no Google, never run: stay quiet
        self.assertEqual(self.runs, [])
        self.assertEqual(jobs.next_run_text(config.load()), "after you run it for the first time")
        self.connect_google()
        self.assertEqual(sched.tick(self.at(self.D5, 8)), "started")

    def test_runs_once_a_day_after_the_daily_time(self):
        self.user_client()
        self.connect_google()
        sched = Scheduler(self.app.manager)
        config.update(lambda c: c.update(schedule={"hour": 7, "minute": 0}))
        self.assertEqual(sched.tick(self.at(self.D5, 6, 59)), "waiting")
        self.assertEqual(sched.tick(self.at(self.D5, 7)), "started")
        self.wait_manager()
        self.assertEqual(sched.tick(self.at(self.D5, 9)), "done today")
        self.assertEqual(self.runs, ["schedule"])
        self.assertEqual(sched.tick(self.at(self.D6, 7, 30)), "started")                # next day: runs again (catch-up too)

    def test_a_failed_run_does_not_use_up_the_day(self):
        def boom(cfg, **kw):
            raise RuntimeError("no network")
        app = self.make_app(boom)
        c = self.user_client(app=app)
        self.connect_google()
        sched = Scheduler(app.manager)
        self.assertEqual(sched.tick(self.at(self.D5, 8)), "started")
        for _ in range(100):
            if not app.manager.snapshot()["running"]:
                break
            time.sleep(0.05)
        self.assertEqual(appdb.meta_get("last_scheduled"), "")                          # not marked done
        self.assertGreater(sched._retry_at, time.time())                                # but waits before trying again
        sched._retry_at = 0
        self.assertEqual(sched.tick(self.at(self.D5, 8, 20)), "started")

    def test_a_run_by_hand_after_the_daily_time_counts_for_the_day(self):
        self.user_client()
        self.connect_google()
        Scheduler(self.app.manager)
        config.update(lambda c: c.update(schedule={"hour": 0, "minute": 0}))        # the daily time has passed, whatever the hour
        self.assertEqual(self.app.manager.start("manual", admin=True)[0], True)
        self.wait_manager()
        self.assertEqual(appdb.meta_get("last_scheduled"), date.today().isoformat())

    def test_a_run_by_hand_before_the_daily_time_does_not(self):
        self.user_client()
        self.connect_google()
        Scheduler(self.app.manager)
        config.update(lambda c: c.update(schedule={"hour": 23, "minute": 59}))
        with mock.patch.object(jobs, "datetime", wraps=datetime) as fake:
            fake.now.return_value = datetime.now().replace(hour=6, minute=0)
            self.assertEqual(self.app.manager.start("manual", admin=True)[0], True)
            self.wait_manager()
        self.assertEqual(appdb.meta_get("last_scheduled"), "")

    def test_sending_the_waiting_companies_does_not_count_as_the_daily_run(self):
        self.user_client()
        self.connect_google()
        sched = Scheduler(self.app.manager)
        sent = []
        with mock.patch.object(self.app.manager, "_sender", lambda cfg, **kw: sent.append(1) or {"status": "ok", "errors": []}):
            self.app.manager.start("connect", admin=True, task="send")
            self.wait_manager()
        self.assertEqual(sent, [1])
        self.assertEqual(appdb.meta_get("last_scheduled"), "")

    def test_the_next_run_text_follows_the_state(self):
        self.user_client()
        self.connect_google()
        cfg = config.load()
        cfg["schedule"] = {"hour": 7, "minute": 0}
        self.assertEqual(jobs.next_run_text(cfg, datetime(2099, 1, 5, 6, 0)), "today 07:00")
        self.assertEqual(jobs.next_run_text(cfg, datetime(2099, 1, 5, 8, 0)), "now (catching up)")
        self.assertEqual(jobs.next_run_text(cfg, datetime(2099, 1, 5, 8, 0), assume_done=True), "tomorrow 07:00")
        self.assertTrue(jobs.counts_for_today(cfg, "schedule", datetime(2099, 1, 5, 3, 0)))
        self.assertTrue(jobs.counts_for_today(cfg, "manual", datetime(2099, 1, 5, 8, 0)))
        self.assertFalse(jobs.counts_for_today(cfg, "manual", datetime(2099, 1, 5, 6, 0)))      # before the daily time
        self.assertFalse(jobs.counts_for_today(cfg, "connect", datetime(2099, 1, 5, 8, 0)))     # a send is not the daily run
        appdb.meta_set("last_scheduled", "2099-01-05")
        self.assertEqual(jobs.next_run_text(cfg, datetime(2099, 1, 5, 8, 0)), "tomorrow 07:00")

    def waiting_row(self):
        self.seed()                                                        # a qualified company that never reached the sheet

    def test_companies_waiting_for_google_are_sent_when_nothing_else_is_going_on(self):
        self.user_client()
        self.connect_google()
        self.waiting_row()
        sent = []
        sched = Scheduler(self.app.manager)
        appdb.meta_set("last_scheduled", "2099-01-05")                      # today's run is done
        with mock.patch.object(self.app.manager, "_sender",
                               lambda cfg, **kw: sent.append(kw["trigger"]) or {"status": "ok", "errors": []}):
            self.assertEqual(sched.tick(self.at(self.D5, 12)), "sending")
            self.wait_manager()
            self.assertEqual(sent, ["connect"])
            self.assertEqual(sched.tick(self.at(self.D5, 12, 1)), "done today")   # not again within ten minutes
            sched._resend_at = 0
            self.assertEqual(sched.tick(self.at(self.D5, 12, 20)), "sending")      # still waiting: the next try
            self.wait_manager()

    def test_nothing_is_sent_without_google_or_without_waiting_companies(self):
        self.user_client()
        sched = Scheduler(self.app.manager)
        appdb.meta_set("last_scheduled", "2099-01-05")
        self.assertEqual(sched.tick(self.at(self.D5, 12)), "not ready")        # no Google, never run
        self.connect_google()
        self.assertEqual(sched.tick(self.at(self.D5, 12)), "done today")       # nothing waiting

    def test_a_sheet_failure_in_a_run_schedules_a_retry_of_the_send(self):
        self.user_client()
        sched = Scheduler(self.app.manager)
        sched._finished("schedule", "partial", "", {"errors": ["sheet: Could not reach Google."]})
        self.assertGreater(sched._resend_at, time.time() + 500)
        sched._resend_at = 0
        sched._finished("schedule", "ok", "", {"errors": []})
        self.assertEqual(sched._resend_at, 0)

    def test_busy_run_is_retried_on_the_next_tick(self):
        self.user_client()
        self.connect_google()
        self.gate.clear()
        self.app.manager.start("manual", admin=True)
        sched = Scheduler(self.app.manager)
        self.assertEqual(sched.tick(self.at(self.D5, 12)), "busy")
        self.gate.set()
        self.wait_manager()
        self.assertEqual(sched.tick(self.at(self.D5, 12, 1)), "started")

    def test_an_unreadable_config_does_not_end_the_schedule(self):
        self.user_client()
        sched = Scheduler(self.app.manager)
        config.CONFIG_PATH.write_text("{ not json")
        with self.assertRaises(config.ConfigError):
            sched.tick(self.at(self.D5, 8))                                              # tick reports it...
        config.CONFIG_PATH.write_text("{}")                                              # ...and the loop carries on
        self.assertEqual(sched.tick(self.at(self.D5, 8)), "not ready")

    def test_a_damaged_settings_file_shows_a_page_not_a_crash_and_never_sticks_a_run(self):
        c = self.user_client()
        config.CONFIG_PATH.write_text("{ not json")
        r = c.get("/")
        self.assertEqual(r.status_code, 500)
        self.assertIn("settings file is damaged", r.get_data(as_text=True))
        config.CONFIG_PATH.write_text("{}")
        self.assertEqual(c.get("/").status_code, 200)
        # a file that breaks only after the run has started must still clear "running"
        real = config.load
        calls = {"n": 0}
        def flaky():
            calls["n"] += 1
            if calls["n"] == 2:                                                          # start() reads it once, the worker again
                raise config.ConfigError("broken")
            return real()
        with mock.patch.object(config, "load", flaky):
            self.assertTrue(self.app.manager.start("manual", admin=True)[0])
            for _ in range(100):
                if not self.app.manager.snapshot()["running"]:
                    break
                time.sleep(0.02)
        snap = self.app.manager.snapshot()
        self.assertFalse(snap["running"])
        self.assertIn("broken", snap["last"]["error"])


class ReadOnlyCopyTests(WebCase):
    """The published copy (Vercel): the list can be looked at and downloaded, nothing else."""

    def test_a_read_only_copy_shows_the_list_but_cannot_run_or_change_settings(self):
        with mock.patch.dict("os.environ", {"ITLEADS_READ_ONLY": "1", "ITLEADS_SNAPSHOT_AT": "6 Oct 2026, 18:45 PKT"}):
            app = self.make_app(self.fake_runner)
            c = self.user_client(app=app)
            self.seed()
            home = c.get("/").get_data(as_text=True)
            self.assertIn("read-only copy", home.lower())
            self.assertIn("6 Oct 2026, 18:45 PKT", home)
            self.assertNotIn('href="/settings"', home)                              # no Settings in the menu
            self.assertRegex(home, r'id="run"[^>]*\bdisabled\b')                    # Run now cannot be pressed
            self.assertNotIn("Connect Google", home)
            r = self.api(c, "/api/run", {})
            self.assertEqual(r.status_code, 403)
            self.assertFalse(r.get_json()["ok"])
            self.assertEqual(self.runs, [])                                         # nothing was started
            for path in ("/api/settings/rules", "/api/settings/schedule", "/api/settings/google", "/api/settings/sharing"):
                self.assertEqual(self.api(c, path, {}).status_code, 403, path)
            self.assertEqual(c.get("/settings").status_code, 403)
            self.assertEqual(c.get("/api/preview").get_json()["total"], 1)           # looking and downloading still work
            self.assertEqual(c.get("/download.csv").status_code, 200)
            self.assertEqual(c.get("/download.xlsx").status_code, 200)
            self.assertEqual(c.get("/api/status").get_json()["next_run"], "")        # it never says when it will run
            t = self.token(c, "/")
            self.assertEqual(c.post("/logout", data={"_csrf": t}).status_code, 302)  # signing out is still allowed

    def test_a_read_only_copy_counts_days_from_the_snapshot_not_from_today(self):
        self.seed()
        old = (date.today() - timedelta(days=40)).isoformat()
        s = Store(config.DATA / "leads.db")
        s.db.execute("UPDATE companies SET ready_at=?, first_seen=?", (old, old))
        s.db.commit()
        s.close()
        c = self.user_client()
        text = lambda: re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", c.get("/").get_data(as_text=True)))
        self.assertRegex(text(), r"Last 7 days 0\b")                                # a live app: nothing was added this week
        with mock.patch.dict("os.environ", {"ITLEADS_READ_ONLY": "1", "ITLEADS_SNAPSHOT_DATE": old}):
            self.assertRegex(text(), r"Last 7 days 1\b")                            # a snapshot: the week it was taken

    def test_the_session_secret_can_come_from_the_environment_and_is_shared_by_every_instance(self):
        secret = "s" * 48
        with mock.patch.dict("os.environ", {"ITLEADS_SECRET_KEY": secret}):
            first, second = self.make_app(self.fake_runner), self.make_app(self.fake_runner)
            self.assertEqual(first.config["SECRET_KEY"], secret)
            self.assertNotEqual(config.load()["app"]["secret_key"], secret)          # the environment's secret is never written to disk
            c = self.user_client(app=first)
            cookie = c.get_cookie("itleads").value
            other = second.test_client()                                            # another serverless instance
            other.set_cookie("itleads", cookie)
            self.assertEqual(other.get("/api/status").status_code, 200)              # the same sign-in works there

    def test_behind_a_proxy_the_visitors_forwarded_address_is_the_one_that_counts(self):
        self.make_user()
        for flag, expect, other in (("1", "203.0.113.9", "10.1.1.1"), ("", "10.1.1.1", "203.0.113.9")):
            with mock.patch.dict("os.environ", {"ITLEADS_PROXY_FIX": flag}):
                app = self.make_app(self.fake_runner)
                c = app.test_client()
                t = self.token(c, "/login")
                c.post("/login", data={"_csrf": t, "email": "admin@example.com", "password": "wrong-one"},
                       headers={"X-Forwarded-For": "203.0.113.9"}, environ_overrides={"REMOTE_ADDR": "10.1.1.1"})
                self.assertEqual(appdb.failures(f"ip:{expect}"), 1, flag)
                self.assertEqual(appdb.failures(f"ip:{other}"), 0, flag)
                appdb.clear_all_lockouts()


class CliTests(WebCase):
    def run_cli(self, *argv):
        import contextlib
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf), mock.patch("os._exit", side_effect=SystemExit):
            try:
                code = cli.main(list(argv))
            except SystemExit as e:                                  # main() ends the process for some exit codes
                code = e.code if e.code is not None else 0
        return code, buf.getvalue()

    def test_the_protected_folder_check_ignores_case_and_lookalikes(self):
        home = Path.home()
        for path, want in ((home / "Desktop" / "it-leads", True), (home / "desktop" / "it-leads", True),
                           (home / "DOCUMENTS", True), (home / "Library" / "Mobile Documents" / "x", True),
                           (home / "it-leads", False), (home / "DocumentsOld" / "it-leads", False),
                           (home / "Desktop-old", False)):
            with mock.patch.object(config, "ROOT", path):
                self.assertEqual(cli.in_protected_folder(), want, str(path))

    def test_the_held_back_export_cannot_run_as_a_formula(self):
        store = Store(config.DATA / "leads.db")
        rec = {"id": "tx:+", "source": "tx", "name": "+Plus Tech LLC", "trade_name": "", "state": "TX", "city": "Austin",
               "address": "=1+1 Main St", "zip": "78701", "registered": "2026-10-02", "industry": "Software",
               "individual": False, "emails": [], "people": []}
        store.add(rec, date.today())
        store.record_attempt("tx:+", {"website": "https://plus.io", "email": "", "phone": "", "domain": "plus.io", "missing": ["email"],
                                      "it": {"level": "strong"}, "why": [], "contact": {}}, "held", date.today())
        store.close()
        out = config.ROOT / "held.csv"
        with mock.patch.dict("os.environ", {"ITLEADS_CWD": str(config.ROOT)}):
            code, text = self.run_cli("export", "--incomplete", "--out", "held.csv")
        self.assertEqual(code, 0)
        body = out.read_text(encoding="utf-8-sig")
        self.assertIn("'+Plus Tech LLC", body)
        self.assertIn("'=1+1 Main St", body)
        self.assertIn(str(out), text)                                  # the message shows the full path

    def test_reset_password_for_the_team_login_lifts_the_lockout_and_changes_nothing(self):
        cfg = config.load()
        auth.ensure_internal_login(cfg)
        team = auth.internal_login(cfg)
        key = appdb.email_key(team["email"])
        before = appdb.user_by_email(team["email"])["pw_hash"]
        for k in (f"acct:{key}", f"u:10.0.0.1|{key}", "ip:10.0.0.1"):
            for _ in range(3):
                appdb.record_failure(k)
        with mock.patch("getpass.getpass", side_effect=AssertionError("must not ask for a password")):
            code, text = self.run_cli("reset-password", team["email"].upper())
        self.assertEqual(code, 0)
        self.assertIn("unlocked", text)
        self.assertIn("config.json", text)                              # says where the real password lives
        for k in (f"acct:{key}", f"u:10.0.0.1|{key}", "ip:10.0.0.1"):
            self.assertEqual(appdb.failures(k), 0, k)
        self.assertEqual(appdb.user_by_email(team["email"])["pw_hash"], before)

    def test_set_login_stores_the_team_login_in_config_and_it_works(self):
        code, text = self.run_cli("set-login", "crew@agency.test", "--generate")
        self.assertEqual(code, 0)
        lg = auth.internal_login(config.load())
        self.assertEqual(lg["email"], "crew@agency.test")
        self.assertIn(lg["password"], text)                                         # shown once, so it can be handed over
        self.assertRegex(lg["password"], r"^[A-Za-z2-9]{4}(-[A-Za-z2-9]{4}){3}$")
        first = lg["password"]
        self.assertEqual(self.run_cli("set-login", "--generate")[0], 0)
        again = auth.internal_login(config.load())
        self.assertNotEqual(again["password"], first)                              # a new one each time, the email is kept
        self.assertEqual(again["email"], "crew@agency.test")
        auth.ensure_internal_login(config.load())
        self.assertEqual(self.login(self.app.test_client(), "crew@agency.test", again["password"]).status_code, 302)

    def test_set_login_refuses_a_short_or_mismatched_password_and_a_bad_email(self):
        before = auth.internal_login(config.load())
        with mock.patch("getpass.getpass", side_effect=["short", "short"]):
            code, text = self.run_cli("set-login")
        self.assertEqual(code, 1)
        self.assertIn("at least", text)
        with mock.patch("getpass.getpass", side_effect=["a-long-enough-password-1", "a-different-password-22"]):
            code, text = self.run_cli("set-login")
        self.assertEqual((code, "differ" in text), (1, True))
        self.assertEqual(self.run_cli("set-login", "not-an-email", "--generate")[0], 1)
        self.assertEqual(auth.internal_login(config.load()), before)               # nothing was changed
        with mock.patch("getpass.getpass", side_effect=["a-long-enough-password-1", "a-long-enough-password-1"]):
            self.assertEqual(self.run_cli("set-login")[0], 0)
        self.assertEqual(auth.internal_login(config.load())["password"], "a-long-enough-password-1")

    def test_reset_password_for_somebody_else_names_the_only_login(self):
        code, text = self.run_cli("reset-password", "nobody@example.com")
        self.assertEqual(code, 1)
        self.assertIn(auth.internal_login(config.load())["email"], text)
        self.assertNotIn("create the first one", text)                  # there is no sign-up any more

    def test_open_says_how_to_start_the_app_when_it_is_not_running(self):
        with mock.patch.object(cli, "wait_until_up", return_value=False), mock.patch("webbrowser.open") as opened:
            code, text = self.run_cli("open")
        self.assertEqual(code, 1)
        opened.assert_not_called()
        self.assertIn("serve --open", text)

    def test_a_broken_environment_or_database_gives_a_plain_message_not_a_traceback(self):
        import sqlite3
        with mock.patch.object(cli, "cmd_status", side_effect=ImportError("No module named 'waitress'")):
            code, text = self.run_cli("status")
        self.assertEqual(code, 1)
        self.assertIn("rm -rf .venv", text)
        with mock.patch.object(cli, "cmd_status", side_effect=sqlite3.DatabaseError("file is not a database")):
            code, text = self.run_cli("status")
        self.assertEqual(code, 1)
        self.assertIn("could not be opened", text)

    def test_serve_on_a_busy_port_says_so_instead_of_pretending(self):
        import argparse
        import socket
        busy = socket.socket()
        busy.bind(("127.0.0.1", 0))
        busy.listen(1)
        port = busy.getsockname()[1]
        try:
            import contextlib
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf), mock.patch.object(cli, "wait_until_up", return_value=False):
                code = cli.cmd_serve(argparse.Namespace(host="127.0.0.1", port=port, open=False))
            self.assertEqual(code, 1)
            self.assertIn("Another program is using that port", buf.getvalue())
            self.assertNotIn("Running at", buf.getvalue())             # no banner for a server that never started
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf), mock.patch.object(cli, "wait_until_up", return_value=True):
                code = cli.cmd_serve(argparse.Namespace(host="127.0.0.1", port=port, open=False))
            self.assertEqual(code, 0)                                   # already running: not a failure (no restart loop)
            self.assertIn("already running", buf.getvalue())
        finally:
            busy.close()

    def test_installing_the_service_where_there_is_no_launchd_explains_it(self):
        with mock.patch("shutil.which", return_value=None):
            code, text = self.run_cli("install-service")
            self.assertEqual(code, 1)
            self.assertIn("macOS", text)
            self.assertEqual(self.run_cli("uninstall-service")[0], 0)


class PipelineStateTests(WebCase):
    def test_an_offline_attempt_leaves_no_row_in_the_history(self):
        from itleads import pipeline
        with mock.patch.object(pipeline, "online", return_value=False):
            summary = pipeline.run(config.load(), quiet=True)
        self.assertEqual(summary["status"], "offline")
        store = Store(config.DATA / "leads.db")
        self.assertIsNone(store.last_run())
        store.close()

    def test_a_push_that_fails_halfway_reports_what_arrived(self):
        from itleads import pipeline
        store = Store(config.DATA / "leads.db")
        for i in range(100):
            rec = {"id": f"tx:{i}", "source": "tx", "name": f"Co {i} LLC", "state": "TX", "city": "Austin", "address": "1 Main St",
                   "zip": "78701", "registered": "2026-10-02", "industry": "Software", "individual": False, "emails": [], "people": []}
            store.add(rec, date.today())
            store.record_attempt(rec["id"], {"website": f"https://co{i}.io", "email": f"hi@co{i}.io", "phone": "5122100100",
                                             "domain": f"co{i}.io", "why": [], "contact": {}, "it": {"level": "strong"}}, "ready", date.today())
        class Flaky:
            calls = 0
            def upsert(self, rows, today):
                Flaky.calls += 1
                if Flaky.calls == 2:
                    raise sheet.BridgeError("Google went away")
                return {"inserted": len(rows), "updated": 0, "skipped": 0}
        with self.assertRaises(sheet.BridgeError) as cm:
            pipeline.push(store, Flaky(), date.today(), lambda *a: None)
        self.assertEqual(cm.exception.partial["pushed"], 80)                       # the first batch really is in the sheet
        self.assertEqual(store.counts()["pushed"], 80)
        self.assertEqual(store.counts()["ready"], 20)                              # the rest waits
        store.close()


class ExportUnitTests(unittest.TestCase):
    EVIL = {"company": "=1+1", "website": "javascript:alert(1)", "email": "+cmd|x@evil.io", "phone": "@SUM(1)",
            "address": "-2+3", "registered": "2026-10-01", "contact": "\x00bad\x07name", "contact_email": "",
            "linkedin": '=HYPERLINK("http://evil.example")', "industry": "", "state": "TX", "source": "tx", "fit": "",
            "proof": "\t=1+1"}

    def test_empty_exports_are_valid(self):
        ws = load_workbook(io.BytesIO(export.to_xlsx([]))).active
        self.assertEqual(ws.max_row, 1)
        self.assertTrue(export.to_csv([]).startswith("﻿Company".encode()))

    def test_nothing_from_a_website_becomes_a_formula_or_a_dangerous_link(self):
        ws = load_workbook(io.BytesIO(export.to_xlsx([self.EVIL]))).active
        cells = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c) for c in range(1, 15)}
        for name in ("COMPANY", "WEBSITE", "EMAIL", "PHONE", "ADDRESS", "LINKEDIN", "VERIFIED BY"):
            self.assertEqual(cells[name].data_type, "s", name)                           # text, never a formula
        self.assertIsNone(cells["WEBSITE"].hyperlink)                                    # javascript: is not a link
        self.assertIsNone(cells["LINKEDIN"].hyperlink)
        self.assertEqual(cells["CONTACT"].value, "badname")                              # control characters are dropped

    def test_characters_a_spreadsheet_file_cannot_hold_never_break_the_download(self):
        bad = dict(self.EVIL, email="ab\uffff@acme.test", contact="Ada\udc00 \ufffe Lovelace", proof="ok\uffff")
        ws = load_workbook(io.BytesIO(export.to_xlsx([bad]))).active              # used to raise 'All strings must be XML compatible'
        cells = {ws.cell(row=1, column=c).value: ws.cell(row=2, column=c).value for c in range(1, 15)}
        self.assertEqual(cells["EMAIL"], "ab@acme.test")
        self.assertIn("Lovelace", cells["CONTACT"])
        self.assertNotIn("\ufffe", cells["CONTACT"])
        export.to_csv([bad])

    def test_csv_keeps_formula_looking_text_inert_and_drops_control_characters(self):
        text = export.to_csv([self.EVIL]).decode("utf-8")
        for want in ("'=1+1", "'+cmd|x@evil.io", "'@SUM(1)", "'-2+3", "badname"):
            self.assertIn(want, text)
        self.assertNotIn("\x00", text)
        self.assertNotIn("\x07", text)


if __name__ == "__main__":
    unittest.main()
